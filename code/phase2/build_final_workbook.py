"""Build results/results_sheet_final.xlsx (one tab per manuscript item) from results/results_sheet.csv. Every value is read from the sheet; each row lists the result IDs of its values."""

import datetime as dt
import hashlib
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
SHEET = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "results/results_sheet.csv"
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "results/results_sheet_final.xlsx"
S = pd.read_csv(SHEET)
S["variant"] = S.variant.fillna("")
S["detail"] = S.detail.fillna("")
ROOMS = ["room2", "room6", "room7"]
TABS = []


def get(section, metric=None, variant=None, fold=None, detail=None):
    d = S[S.section == section]
    for col, val in (
        ("metric", metric),
        ("variant", variant),
        ("room_or_fold", fold),
        ("detail", detail),
    ):
        if val is not None:
            d = d[d[col] == val]
    assert len(d) == 1, (section, metric, variant, fold, detail, len(d))
    return d.iloc[0]


def row(*cells):
    """cells: label strings or sheet rows; returns (values, ids)."""
    vals, ids = [], []
    for c in cells:
        if isinstance(c, pd.Series):
            vals.append(c.value)
            ids.append(c.result_id)
        else:
            vals.append(c)
    return vals, ids


def table(columns, rows, fmt=None, id_col="Result IDs"):
    recs = []
    for cells in rows:
        vals, ids = row(*cells)
        recs.append(vals + ["; ".join(ids)])
    return pd.DataFrame(recs, columns=columns + [id_col]), fmt or {}


def add_tab(name, blocks):
    """blocks: list of (title, DataFrame, {column: number_format})."""
    TABS.append((name, blocks))


def ci_text(r, dec=2):
    return f"{r.value:.{dec}f} [{r.ci_lo:.{dec}f}-{r.ci_hi:.{dec}f}]"


F2, F3, F0 = "0.00", "0.000", "#,##0"
metrics_cols = ["MAE (days)", "RMSE (days)", "R2"]


def metric_row(section, variant, fold="macro"):
    return [get(section, m, variant, fold) for m in ("MAE", "RMSE", "R2")]


COVERAGE = [
    ("Audio", "observed hours: audio"),
    ("Tracking (YOLO)", "observed hours: tracking (YOLO)"),
    ("Optical flow", "observed hours: optical flow"),
    ("Morphometry before QC", "observed hours: morphometry (before QC filter)"),
    ("Morphometry after QC", "observed hours: morphometry (after QC filter)"),
    (
        "Complete case (primary)",
        "complete-case hours, primary (audio + tracking + morphometry after QC)",
    ),
]
DENOMINATOR = "denominator: hours from placement to last observed hour"
rows = []
for room in ROOMS + ["all rooms"]:
    total = get(
        "S01_data_coverage", "hours from placement to the room's last observed hour", None, room
    )
    for label, metric in COVERAGE:
        hours = get("S01_data_coverage", metric, None, room, DENOMINATOR)
        rows.append([room, label, hours, total, round(100 * hours.value / total.value, 1)])
df, _ = table(
    [
        "Room",
        "Modality",
        "Observed hours",
        "Hours from placement to last observed hour",
        "% of hours",
    ],
    rows,
)
add_tab(
    "Table1_coverage",
    [
        (
            "Table 1. Observed hours per room and modality (denominator: hours from placement, 5 June 2025, to the room's last observed hour)",
            df,
            {
                "Observed hours": F0,
                "Hours from placement to last observed hour": F0,
                "% of hours": "0.0",
            },
        )
    ],
)

rows = [
    ["Fused multimodal (HGB)"] + metric_row("S30_loro_feature_sets", "Fused"),
    ["Video only (morphometry + motion + flow)"] + metric_row("S30_loro_feature_sets", "All video"),
    ["Morphometry only"] + metric_row("S30_loro_feature_sets", "Morphometry"),
    ["Audio only (bioacoustics + acoustic level)"]
    + metric_row("S34_fusion_by_learner", "HistGradientBoosting | Audio only"),
    ["Baseline: median body height (isotonic)"]
    + metric_row("S39_baselines", "median body height (isotonic)"),
    ["Baseline: global mean"] + metric_row("S39_baselines", "global mean"),
]
df, _ = table(["Model"] + metrics_cols, rows)
add_tab(
    "Table2_LORO",
    [
        (
            "Table 2. Leave-one-room-out performance (mean of the three held-out rooms; 4,198 complete-case 6 h windows; the median-body-height baseline is scored on the windows that have the size feature)",
            df,
            {"MAE (days)": F2, "RMSE (days)": F2, "R2": F3},
        )
    ],
)

BANDS = ["<45d", "45-75d", "75-105d", "105-140d", ">140d"]
BLK = [
    ("morph", "Morphometry"),
    ("audio_bio", "Bioacoustics"),
    ("yolo", "Spatial behaviour"),
    ("flow", "Optical flow"),
    ("audio_equip", "Acoustic level"),
    ("time/flags", "Time and data quality"),
]
for tab, sec, met, title, fm in (
    (
        "Table3_SHAP_share",
        "S21_attribution",
        "share of total |SHAP|",
        "Table 3. Modality share of SHAP attribution by age band (mean of the three held-out rooms)",
        F3,
    ),
    (
        "Table4_SHAP_abs",
        "S37_attribution_abs",
        "mean |SHAP| (days)",
        "Table 4. Absolute modality attribution by age band (mean |SHAP|, days; mean of the three held-out rooms)",
        F2,
    ),
):
    rows = [[b] + [get(sec, met, f"primary | {k}", b) for k, _ in BLK] for b in BANDS]
    df, _ = table(["Age band"] + [n for _, n in BLK], rows)
    add_tab(tab, [(title, df, {n: fm for _, n in BLK})])

FS5 = [
    ("Fused", "Fused"),
    ("Video only", "All video"),
    ("Morphometry only", "Morphometry"),
    ("Audio only (bioacoustics only)", "Bioacoustics"),
]
rows = [
    [
        a,
        get("S31_phase_mae_feature_sets", "MAE growth (<=70 d)", b, "macro"),
        get("S31_phase_mae_feature_sets", "MAE plateau (>70 d)", b, "macro"),
    ]
    for a, b in FS5
]
df, _ = table(["Feature set", "Growth (<=70 d) MAE (days)", "Plateau (>70 d) MAE (days)"], rows)
add_tab(
    "Table5_phase",
    [
        (
            "Table 5. Phase-stratified MAE (days), leave-one-room-out; mean of the three rooms' MAE",
            df,
            {"Growth (<=70 d) MAE (days)": F2, "Plateau (>70 d) MAE (days)": F2},
        )
    ],
)

rows = [
    [r] + [get("S23_clustering", f"silhouette_k{k}", "primary", r) for k in range(2, 7)]
    for r in ROOMS
]
df, _ = table(["Room", "K=2", "K=3", "K=4", "K=5", "K=6"], rows)
add_tab(
    "Table6_silhouette",
    [
        (
            "Table 6. Silhouette score by room and number of clusters",
            df,
            {f"K={k}": F3 for k in range(2, 7)},
        )
    ],
)
rows = [
    [
        r,
        70,
        get("S23_clustering", "primary_transition_day", "primary", r),
        get(
            "S38_clustering_persistence", "persistence criterion, 3 consecutive days (day)", None, r
        ),
        get("S23_clustering", "k2_boundary_day", "primary", r),
    ]
    for r in ROOMS
]
df, _ = table(
    [
        "Room",
        "Supervised boundary (fixed analysis parameter)",
        "50% crossing",
        "Persistence criterion, 3 days",
        "Maximum balanced accuracy (k=2 boundary)",
    ],
    rows,
)
add_tab(
    "Table7_transition",
    [("Table 7. Transition-day estimates (days)", df, {c: F0 for c in df.columns[1:-1]})],
)

rows = []
for k in (5, 10, 15, 20):

    def cell(ph, m, fold="all rooms"):
        return get(
            "S10_detection_ci_incl_calibration",
            m,
            f"all_sensor | {ph} | k={k} d | rule=same_hour",
            fold,
        )

    cells = [k]
    for ph in ("growth (<=70 d)", "plateau (>70 d)"):
        for m in ("AUC", "sensitivity"):
            r = cell(ph, m)
            cells.append(r)
    for ph in ("growth (<=70 d)", "plateau (>70 d)"):
        for room in ROOMS:
            cells.append(cell(ph, "achieved specificity", room))
    rows.append(cells)
cols = [
    "Injected delay (days)",
    "Growth AUC",
    "Growth sensitivity",
    "Plateau AUC",
    "Plateau sensitivity",
] + [f"Achieved specificity, {p}, {r}" for p in ("growth", "plateau") for r in ROOMS]
df, _ = table(cols, rows)
for c in ("Growth AUC", "Growth sensitivity", "Plateau AUC", "Plateau sensitivity"):
    ci = []
    for k in (5, 10, 15, 20):
        ph = "growth (<=70 d)" if c.startswith("Growth") else "plateau (>70 d)"
        m = "AUC" if "AUC" in c else "sensitivity"
        ci.append(
            ci_text(
                get(
                    "S10_detection_ci_incl_calibration",
                    m,
                    f"all_sensor | {ph} | k={k} d | rule=same_hour",
                    "all rooms",
                )
            )
        )
    df[c] = ci
add_tab(
    "Table8_detection",
    [
        (
            "Table 8. Detection of injected developmental delay (all sensors, calibrated detector, same-hour matched anchors, all rooms; estimate [95% CI, room-week bootstrap including recalibration]; sensitivity at the threshold set for 95% specificity on the training rooms)",
            df,
            {c: F2 for c in df.columns if c.startswith("Achieved")},
        )
    ],
)

blocks = ["31-70d", "70-110d", "110-138d", "138-160d", "160-182d"]
rows = [
    [
        i + 1,
        b,
        get("S05_ageblock", "n", "primary fused", b, "AGEBLOCK_12h"),
        get("S05_ageblock", "MAE", "primary fused", b, "AGEBLOCK_12h"),
        get("S05_ageblock", "R2", "primary fused", b, "AGEBLOCK_12h"),
    ]
    for i, b in enumerate(blocks)
]
rows += [
    [
        "pooled fold-mean",
        "",
        "",
        get("S05_ageblock", "MAE", "primary fused", "pooled fold-mean", "AGEBLOCK_12h"),
        "",
    ],
    [
        "pooled fold-mean excl. block 1",
        "",
        "",
        get(
            "S05_ageblock", "MAE", "primary fused", "pooled fold-mean excl. block 1", "AGEBLOCK_12h"
        ),
        "",
    ],
]
df, _ = table(["Block", "Age range (days, rounded block edges)", "n", "MAE (days)", "R2"], rows)
add_tab(
    "Table9_ageblock",
    [
        (
            "Table 9. Leave-one-age-block-out performance by block (12 h window purge, primary fused model)",
            df,
            {"n": F0, "MAE (days)": F2, "R2": F3},
        )
    ],
)

mf = S[
    (S.section == "S32_model_family") & (S.room_or_fold == "macro") & (S.metric == "MAE")
].sort_values("value")
rows = []
for v in mf.variant:
    r = [get("S32_model_family", m, v, "macro") for m in ("MAE", "RMSE", "R2")]
    rows.append([v] + r + [r[0].n, r[0].detail])
df, _ = table(["Model"] + metrics_cols + ["n windows", "Imputation / scaling / notes"], rows)
add_tab(
    "Fig5_model_family",
    [
        (
            "Fig 5. Model families, leave-one-room-out, primary pipeline (mean of the three rooms)",
            df,
            {"MAE (days)": F2, "RMSE (days)": F2, "R2": F3, "n windows": F0},
        )
    ],
)

bs = S[S.section == "S36_body_size"]
df = pd.DataFrame(
    {
        "Room / scope": bs.room_or_fold,
        "Metric": bs.metric,
        "Value": bs.value,
        "Result ID": bs.result_id,
    }
)
add_tab(
    "Fig9_body_size",
    [
        (
            "Fig 9. Body-size trajectory: relative body height as a ratio to each room's first-week baseline",
            df,
            {"Value": F2},
        )
    ],
)

rows = []
for c in S[S.section == "S07_paired_fusion"].variant.unique():
    for fold in ROOMS + ["macro"]:
        r = get("S07_paired_fusion", "MAE difference (days)", c, fold)
        rows.append([c, fold, r, r.ci_lo, r.ci_hi, r.n])
df, _ = table(
    [
        "Contrast (first minus second)",
        "Room",
        "MAE difference (days)",
        "95% CI low",
        "95% CI high",
        "n windows",
    ],
    rows,
)
add_tab(
    "S1_paired_fusion",
    [
        (
            "Fig S1. Paired MAE differences on identical held-out windows (negative: the first model has the lower error; room-day block bootstrap, 2,000 resamples)",
            df,
            {"MAE difference (days)": F2, "95% CI low": F2, "95% CI high": F2, "n windows": F0},
        )
    ],
)

lrn = []
for L in ("ExtraTrees", "HistGradientBoosting"):
    for fs in ("Fused", "Video only", "Morphometry only", "Audio only"):
        lrn.append(
            [L, fs]
            + [
                get("S34_fusion_by_learner", m, f"{L} | {fs}", "macro")
                for m in ("MAE", "RMSE", "R2")
            ]
            + [get("S34_fusion_by_learner", "MAE", f"{L} | {fs}", r) for r in ROOMS]
        )
d1, _ = table(
    ["Learner", "Feature set"] + metrics_cols + ["MAE room2", "MAE room6", "MAE room7"], lrn
)
rows = []
for c in S[S.section == "S35_fusion_paired_by_learner"].variant.unique():
    for fold in ROOMS + ["macro"]:
        r = get("S35_fusion_paired_by_learner", "MAE difference (days)", c, fold)
        rows.append([c, fold, r, r.ci_lo, r.ci_hi, r.n])
d2, _ = table(
    ["Contrast", "Room", "MAE difference (days)", "95% CI low", "95% CI high", "n windows"], rows
)
rows = []
for sub in (
    "(a) complete-case windows",
    "(b) all eligible windows",
    "(c) hours without video only",
):
    for L in ("HistGradientBoosting", "ExtraTrees"):
        v = f"{L} | {sub}"
        rows.append(
            [
                sub,
                L,
                get("S33_missing_data", "MAE", v, "macro"),
                get("S33_missing_data", "n", v, "room2"),
                get("S33_missing_data", "n", v, "room6"),
                get("S33_missing_data", "n", v, "room7"),
            ]
            + [get("S33_missing_data", "MAE", v, r) for r in ROOMS]
        )
d3, _ = table(
    [
        "Subset",
        "Learner",
        "Macro MAE (days)",
        "n room2",
        "n room6",
        "n room7",
        "MAE room2",
        "MAE room6",
        "MAE room7",
    ],
    rows,
)
fmt2 = {
    c: F2
    for c in (
        "MAE (days)",
        "RMSE (days)",
        "MAE room2",
        "MAE room6",
        "MAE room7",
        "MAE difference (days)",
        "95% CI low",
        "95% CI high",
        "Macro MAE (days)",
    )
}
add_tab(
    "ExtraTrees_vs_HGB",
    [
        ("Fusion under ExtraTrees and HGB (LORO macro; four feature sets)", d1, {**fmt2, "R2": F3}),
        (
            "Paired MAE differences (room-day bootstrap 95% CI; negative: first model better)",
            d2,
            {**fmt2, "n windows": F0},
        ),
        (
            "HGB vs ExtraTrees with missing data (macro MAE; n = windows per room)",
            d3,
            {**fmt2, "n room2": F0, "n room6": F0, "n room7": F0},
        ),
    ],
)

inv = S[
    (S.section == "S17_regime_inventory")
    & (S.room_or_fold == "all rooms")
    & (S.metric == "anchors_total")
]
rows = []
for v in inv.variant:
    a = get("S17_regime_inventory", "anchors_total", v, "all rooms")
    cells = [v, a]
    for mdl in (
        "A: primary complete-case fused",
        "B: fused trained on all hours",
        "C: regime-specific (present blocks only)",
    ):
        d_ = S[
            (S.section == "S18_regime_results")
            & (S.metric == "MAE_macro")
            & (S.variant == f"{v} | {mdl}")
        ]
        cells.append(d_.iloc[0] if len(d_) else "not scored (fewer than 30 anchors or one room)")
    rows.append(cells)
df, _ = table(
    [
        "Regime (blocks present)",
        "Anchors (6 h windows)",
        "MAE A: primary model",
        "MAE B: trained on all hours",
        "MAE C: regime-specific",
    ],
    rows,
)
add_tab(
    "Availability_regimes",
    [
        (
            "Availability regimes: macro MAE (days) over the rooms with at least 30 anchors in the regime",
            df,
            {
                "Anchors (6 h windows)": F0,
                "MAE A: primary model": F2,
                "MAE B: trained on all hours": F2,
                "MAE C: regime-specific": F2,
            },
        )
    ],
)

rows = []
for scope in ROOMS + ["pooled (all held-out windows)", "macro (mean of rooms)"]:
    rows.append(
        [scope]
        + [
            get("S19_classification", m, None, scope)
            for m in ("n", "accuracy", "balanced_accuracy", "macro_F1", "adjacent_accuracy", "qwk")
        ]
    )
d1, _ = table(
    [
        "Scope",
        "n",
        "Accuracy",
        "Balanced accuracy",
        "Macro-F1",
        "Adjacent accuracy",
        "Quadratic weighted kappa",
    ],
    rows,
)
LABS = ["Brooding/early (<60 d)", "Growing (60-110 d)", "Mature (>110 d)"]
rows = []
for scope in ROOMS + ["pooled"]:
    for t in LABS:
        rows.append(
            [scope, t]
            + [get("S20_confusion", f"true {t} -> predicted {p}", None, scope) for p in LABS]
        )
d2, _ = table(["Scope", "True stage"] + [f"Predicted {p}" for p in LABS], rows)
add_tab(
    "Classification",
    [
        (
            "Three-stage classifier, leave-one-room-out",
            d1,
            {"n": F0, **{c: F3 for c in d1.columns[2:-1]}},
        ),
        ("Confusion counts (windows)", d2, {c: F0 for c in d2.columns[2:-1]}),
    ],
)

rows = []
for v in (
    "clock-only (hour of day)",
    "sensor-health-only",
    "global-mean baseline",
    "fused without acoustic level",
):
    rows.append([v] + metric_row("S26_negative_controls", v))
for v in S[S.section == "S29_thermal_controls"].variant.unique():
    rows.append(
        [v]
        + [
            get("S29_thermal_controls", m, v, "mean over folds and seeds")
            for m in ("MAE", "RMSE", "R2")
        ]
    )
for fold in ("train_summer_test_fall", "train_fall_test_summer"):
    rows.append(
        [f"seasonal extrapolation, {fold}"]
        + [get("S06_season", m, "primary fused", fold, "SEASON_12h") for m in ("MAE", "RMSE", "R2")]
    )
sh = [
    get("S14_shuffled_label", f"{m} (shuffled mean, 95% permutation range)", None, "macro")
    for m in ("MAE", "RMSE", "R2")
]
rows.append(["shuffled labels (100 room-day block permutations, mean)"] + sh)
df, _ = table(["Control"] + metrics_cols, rows)
add_tab(
    "Controls",
    [
        (
            "Negative controls (LORO macro; thermal diagnostics are means over folds and seeds)",
            df,
            {"MAE (days)": F2, "RMSE (days)": F2, "R2": F3},
        ),
        (
            "Shuffled-label control: permutation range and p-values",
            *table(
                ["Quantity", "Value", "95% permutation range low", "95% permutation range high"],
                [
                    ["MAE (days)", sh[0], sh[0].ci_lo, sh[0].ci_hi],
                    ["R2", sh[2], sh[2].ci_lo, sh[2].ci_hi],
                    [
                        "p-value MAE",
                        get("S14_shuffled_label", "p-value MAE", None, "macro"),
                        "",
                        "",
                    ],
                ],
            )[:1],
            {"Value": F3, "95% permutation range low": F3, "95% permutation range high": F3},
        ),
    ],
)

d = S[S.section == "S22_daily_demeaning"]
df = pd.DataFrame(
    {
        "Variant": d.variant,
        "Metric": d.metric,
        "Room / band": d.room_or_fold,
        "Value": d.value,
        "n": d.n,
        "Result ID": d.result_id,
    }
)
add_tab(
    "Demeaning",
    [("Daily demeaning comparison (primary and Step 0 settings)", df, {"Value": F3, "n": F0})],
)

d = S[S.section.isin(["S23_clustering", "S38_clustering_persistence"])]
df = pd.DataFrame(
    {
        "Room": d.room_or_fold,
        "Variant": d.variant,
        "Metric": d.metric,
        "Value": d.value,
        "n": d.n,
        "Result ID": d.result_id,
    }
)
add_tab(
    "Clustering",
    [("Unsupervised clustering and transition-day estimates", df, {"Value": F3, "n": F0})],
)

stages = [
    (
        "Step 0 (centred filter, row slope, anchor purge)",
        "S25_cumulative_states",
        "Step 0 (centred filter, row slope, anchor purge) | fused",
        "Step 0 (centred filter, row slope, anchor purge) | morph_only",
    ),
    (
        "Step 1 (past-only filter, row slope, anchor purge)",
        "S25_cumulative_states",
        "Step 1 (past-only filter, row slope, anchor purge) | fused",
        "Step 1 (past-only filter, row slope, anchor purge) | morph_only",
    ),
    (
        "Step 2 (+ 12 h window purge)",
        "S25_cumulative_states",
        "Step 2 (+ 12 h window purge, age-block and seasonal) | fused",
        "Step 2 (+ 12 h window purge, age-block and seasonal) | morph_only",
    ),
]
rows = []
for nm, sec, vf, vm in stages:
    rows.append(
        [
            nm,
            get(sec, "MAE", vf, "macro", "LORO"),
            get(sec, "MAE", vm, "macro", "LORO"),
            get(sec, "MAE", vf, "pooled fold-mean", "AGEBLOCK_12h"),
        ]
    )
rows.append(
    [
        "Step 3 = primary (time-based slopes, 50% coverage)",
        get("S03_loro", "MAE", "primary fused", "macro"),
        get("S03_loro", "MAE", "primary morphometry-only", "macro"),
        get("S05_ageblock", "MAE", "primary fused", "pooled fold-mean", "AGEBLOCK_12h"),
    ]
)
rows.append(
    [
        "V2 (audio referenced to first week)",
        get("S03_loro", "MAE", "V2 fused (audio first-week reference)", "macro"),
        "",
        get(
            "S05_ageblock",
            "MAE",
            "V2 fused (audio first-week reference)",
            "pooled fold-mean",
            "AGEBLOCK_12h",
        ),
    ]
)
rows.append(
    [
        "V3 (body-size-only comparator)",
        get("S03_loro", "MAE", "V3 body-size only", "macro"),
        "",
        get("S05_ageblock", "MAE", "V3 body-size only", "pooled fold-mean", "AGEBLOCK_12h"),
    ]
)
df, _ = table(
    [
        "Analysis state",
        "LORO MAE fused (days)",
        "LORO MAE morphometry-only (days)",
        "Age-block pooled MAE fused (days)",
    ],
    rows,
)
add_tab(
    "Corrections_before_after",
    [("Corrections, Step 0 to V3 (macro MAE, days)", df, {c: F2 for c in df.columns[1:-1]})],
)

full = S.rename(columns={"result_id": "Result ID"})

readme = [
    ("Results workbook", "results/results_sheet_final.xlsx"),
    ("Built", dt.date.today().isoformat()),
    (
        "Source",
        "results/results_sheet.csv is the machine-readable source (one row per number, each with a result ID); this workbook only reformats it.",
    ),
    (
        "Pipeline settings",
        "SCALE_FILTER_MODE=B_exclude (past-only camera-scale filter, burn-in excluded); SLOPE_MODE=time, SLOPE_MIN_FRAC=0.5; 12 h window purge (24 h sensitivity); HistGradientBoosting random_state 0; 4,198 complete-case 6 h windows (1,372 / 1,133 / 1,693 in rooms 2 / 6 / 7).",
    ),
    (
        "Reproduce",
        "python run_all.py   (empty cache; writes the sheet, this workbook's source CSV, tables and manifest). Then: python code/phase2/build_final_workbook.py and python make_figures.py",
    ),
    ("results_sheet.csv sha256", hashlib.sha256(SHEET.read_bytes()).hexdigest()),
    (
        "Hashes and code version",
        "results/results_sheet_manifest.txt (master data, input files, code hash, feature-matrix hashes, package versions)",
    ),
    (
        "Rounding",
        "Days to 2 decimals; R2 and shares to 3; counts as integers. Values are stored rounded as displayed; full precision is in Full_sheet and the CSV.",
    ),
]
tab_notes = {
    "Table1_coverage": "Table 1",
    "Table2_LORO": "Table 2",
    "Table3_SHAP_share": "Table 3",
    "Table4_SHAP_abs": "Table 4",
    "Table5_phase": "Table 5",
    "Table6_silhouette": "Table 6",
    "Table7_transition": "Table 7",
    "Table8_detection": "Table 8 (with achieved specificity per room)",
    "Table9_ageblock": "Table 9",
    "Fig5_model_family": "Fig 5 data",
    "Fig9_body_size": "Fig 9 data",
    "S1_paired_fusion": "Fig S1 data",
    "ExtraTrees_vs_HGB": "Fusion under ExtraTrees vs HGB; HGB vs ExtraTrees with missing data",
    "Availability_regimes": "Modality-availability regimes",
    "Classification": "Three-stage classifier and confusion counts",
    "Controls": "Negative controls",
    "Demeaning": "Daily demeaning",
    "Clustering": "Clustering and transition days",
    "Corrections_before_after": "Step 0 to V3 corrections",
    "Full_sheet": "All rows with result IDs",
}

wb = Workbook()
ws = wb.active
ws.title = "README"
ws.append(["Item", "Content"])
for k, v in readme:
    ws.append([k, v])
ws.append([])
ws.append(["Tab", "Contents"])
for k in list(tab_notes):
    ws.append([k, tab_notes[k]])
ws.column_dimensions["A"].width = 28
ws.column_dimensions["B"].width = 140
for c in ws[1]:
    c.font = Font(bold=True)
ws.freeze_panes = "A2"
for name, blocks in TABS + [("Full_sheet", [("Full sheet", full, {})])]:
    ws = wb.create_sheet(name)
    r0 = 1
    first = True
    for title, df, fmt in blocks:
        if name != "Full_sheet":
            ws.cell(r0, 1, title).font = Font(bold=True)
            r0 += 1
        for j, c in enumerate(df.columns, 1):
            cell = ws.cell(r0, j, c)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        hdr = r0
        for i, rec in enumerate(df.itertuples(index=False), 1):
            for j, (col, v) in enumerate(zip(df.columns, rec), 1):
                if isinstance(v, float) and pd.isna(v):
                    v = None
                if name != "Full_sheet" and isinstance(v, float):
                    dec = {F0: 0, F2: 2, F3: 3}.get(
                        fmt.get(col), 3 if col in ("R2", "Value") else 2
                    )
                    v = round(v, dec)
                cell = ws.cell(hdr + i, j, v)
                if col in fmt:
                    cell.number_format = fmt[col]
        r0 = hdr + len(df) + 3
        if first:
            ws.freeze_panes = ws.cell(hdr + 1, 1)
            first = False
    for j in range(1, ws.max_column + 1):
        width = max(
            (len(str(c.value)) if c.value is not None else 0) for c in ws[get_column_letter(j)][:60]
        )
        ws.column_dimensions[get_column_letter(j)].width = min(max(10, width + 2), 60)
wb.save(OUT)
print("saved", OUT)
